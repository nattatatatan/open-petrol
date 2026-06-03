"""Backfill the daily-lows series from public Data.NSW FuelCheck price-history files.

Why this exists (CLAUDE.md §10): the live FuelCheck API has NO trends endpoint, and
a freshly deployed cache has zero accumulated history — so the price-cycle advisor
would be dead for weeks. The public Data.NSW history files (monthly XLSX, one row per
operator price submission with a `PriceUpdatedDate`) let us seed real, dated daily
lows at launch, then accumulate from polls afterwards.

The XLSX columns are well-established but verify on first download; the column-name
matching here is deliberately fuzzy. The aggregation (events -> daily area lows) is a
pure function so it can be unit-tested without a real file.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from app.models import DailyLow

# Map FuelCheck history fuel labels/codes to our codes where they differ.
_FUEL_ALIASES = {
    "E10": "E10",
    "U91": "U91",
    "P95": "P95",
    "PREMIUM 95": "P95",
    "P98": "P98",
    "PREMIUM 98": "P98",
    "DL": "DL",
    "DIESEL": "DL",
    "PDL": "PDL",
    "PREMIUM DIESEL": "PDL",
    "LPG": "LPG",
}


@dataclass
class PriceEvent:
    date: str       # ISO YYYY-MM-DD
    fuel_type: str
    price: float


def daily_lows_from_events(events: list[PriceEvent], area: str) -> list[DailyLow]:
    """Collapse raw price-submission events to one lowest price per (day, fuel)."""
    best: dict[tuple[str, str], float] = {}
    for e in events:
        key = (e.date, e.fuel_type)
        if key not in best or e.price < best[key]:
            best[key] = e.price
    rows = [
        DailyLow(date=d, fuel_type=f, area=area, price=p)
        for (d, f), p in best.items()
    ]
    rows.sort(key=lambda r: (r.date, r.fuel_type))
    return rows


def _norm_fuel(raw: str) -> str | None:
    return _FUEL_ALIASES.get(str(raw).strip().upper())


def _norm_date(raw) -> str | None:
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(raw).strip(), fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_history_xlsx(path: Path) -> list[PriceEvent]:
    """Read a Data.NSW price-history XLSX into normalised price events.

    Tolerant of column ordering: locates the FuelCode, PriceUpdatedDate and Price
    columns by header name from the first row.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]

    def col(*names: str) -> int | None:
        for n in names:
            if n in header:
                return header.index(n)
        return None

    i_fuel = col("fuelcode", "fuel code", "fueltype", "fuel type")
    i_date = col("priceupdateddate", "price updated date", "date")
    i_price = col("price")
    if None in (i_fuel, i_date, i_price):
        raise ValueError(f"Unexpected history columns: {header}")

    events: list[PriceEvent] = []
    for row in rows:
        fuel = _norm_fuel(row[i_fuel]) if row[i_fuel] is not None else None
        date = _norm_date(row[i_date]) if row[i_date] is not None else None
        if fuel is None or date is None or row[i_price] is None:
            continue
        try:
            price = float(row[i_price])
        except (TypeError, ValueError):
            continue
        events.append(PriceEvent(date=date, fuel_type=fuel, price=price))
    return events
